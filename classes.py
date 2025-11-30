from pathlib import Path
import os

class Folder:
    def __init__(self, name : str, parent : Path):
        self.name = name
        self.parents = parent
        self.fullPath = parent.joinpath(name)
            
    def create(self):
        os.makedirs(self.fullPath, exist_ok=True)
    
    def __repr__(self):
        return f"Folder(Name: {self.name}, Path: {self.fullPath} )"
    
class File:
    def __init__(self, name :str, parent : Path):
        self.name = name.split('.')[0]
        self.extension = ".".join(name.split('.')[1:]) if '.' in name else ''
        self.parents = parent
        self.fullPath = parent.joinpath(name)
    
    def create(self):
        if not self.parents.exists():
            os.makedirs(self.parents, exist_ok=True)

        if self.extension in ['docx', 'doc']:
            import docx
            doc = docx.Document()
            doc.save(self.fullPath)
        elif self.extension in ['xlsx', 'xls']:
            import openpyxl
            excel = openpyxl.Workbook()
            excel.save(self.fullPath)
        elif self.extension in ['pptx', 'ppt']:
            import pptx
            prs = pptx.Presentation()
            prs.save(self.fullPath)
        else:
            Path(self.fullPath).touch(exist_ok=True)
    
    def __repr__(self):
        return f"File(Name: {self.name}, Extension: {self.extension}, Path: {self.fullPath} )"